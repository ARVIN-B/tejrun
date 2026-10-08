from docx import Document
from django.test import SimpleTestCase
from io import BytesIO
from zipfile import ZipFile

from apps.article_agent.infrastructure.document.docx_generator import DocxGenerator


class DocxGeneratorTests(SimpleTestCase):
    def test_generated_docx_contains_title_requested_headings_and_article_parts(self) -> None:
        article = {
            "title": "Article title",
            "sections": [{"heading": "First heading", "content": "First content"}, {"heading": "Second heading", "content": "Second content"}],
            "conclusion": "Conclusion", "FAQ": [{"question": "Question", "answer": "Answer"}],
            "common_mistakes": "Mistakes", "applications": "Applications",
        }
        output = DocxGenerator().generate(article)
        document = Document(output)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for expected in ("Article title", "First heading", "Second heading", "Conclusion", "Question", "Answer", "Mistakes", "Applications"):
            self.assertIn(expected, text)

    def test_generated_docx_is_a_readable_rtl_b_nazanin_package(self) -> None:
        output = DocxGenerator().generate({
            "title": "راهنمای انرژی", "sections": [{"heading": "بخش اول", "content": "متن فارسی معتبر"}],
            "conclusion": "نتیجه", "FAQ": [{"question": "سوال چیست؟", "answer": "پاسخ است."}],
            "common_mistakes": "اشتباه", "applications": "کاربرد",
        })
        self.assertIsNone(ZipFile(BytesIO(output.getvalue())).testzip())
        document = Document(output)
        heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "بخش اول")
        self.assertTrue(heading.runs[0].bold)
        self.assertEqual(heading.runs[0].font.name, "B Nazanin")
        self.assertIn("w:bidi", heading._p.xml)
