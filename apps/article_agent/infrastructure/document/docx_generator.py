from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


class DocxGenerator:
    FONT_NAME = "B Nazanin"
    FONT_SIZE = 14
    TITLE_SIZE = 20
    HEADING_SIZE = 16

    def generate(self, article: dict) -> BytesIO:
        document = Document()

        self._configure_document(document)

        # ============================================
        # Title
        # ============================================

        title = article.get("title", "").strip()

        if title:
            paragraph = document.add_paragraph()

            self._set_paragraph_rtl(paragraph)
            self._set_paragraph_alignment(
                paragraph,
                WD_ALIGN_PARAGRAPH.CENTER,
            )

            run = paragraph.add_run(title)

            self._set_run_font(
                run,
                size=self.TITLE_SIZE,
                bold=True,
            )

        # ============================================
        # Sections
        # ============================================

        for section in article.get("sections", []):
            heading = section.get("heading", "").strip()
            content = section.get("content", "").strip()

            if heading:
                self._add_heading(
                    document,
                    heading,
                )

            if content:
                self._add_content(
                    document,
                    content,
                )

        # ============================================
        # Conclusion
        # ============================================

        conclusion = article.get(
            "conclusion",
            "",
        ).strip()

        if conclusion:
            self._add_heading(
                document,
                "نتیجه‌گیری",
            )

            self._add_content(
                document,
                conclusion,
            )

        # ============================================
        # FAQ
        # ============================================

        faq = article.get("FAQ", [])

        if faq:
            self._add_heading(
                document,
                "سوالات متداول",
            )

            for item in faq:
                question = item.get(
                    "question",
                    "",
                ).strip()

                answer = item.get(
                    "answer",
                    "",
                ).strip()

                if question:
                    self._add_faq_question(
                        document,
                        question,
                    )

                if answer:
                    self._add_content(
                        document,
                        answer,
                    )

        # ============================================
        # Common Mistakes
        # ============================================

        common_mistakes = article.get(
            "common_mistakes",
            "",
        ).strip()

        if common_mistakes:
            self._add_heading(
                document,
                "اشتباهات رایج",
            )

            self._add_content(
                document,
                common_mistakes,
            )

        # ============================================
        # Applications
        # ============================================

        applications = article.get(
            "applications",
            "",
        ).strip()

        if applications:
            self._add_heading(
                document,
                "کاربردها",
            )

            self._add_content(
                document,
                applications,
            )

        # ============================================
        # Save to memory
        # ============================================

        output = BytesIO()

        document.save(output)

        output.seek(0)

        return output

    # ==================================================
    # Document configuration
    # ==================================================

    def _configure_document(
        self,
        document: Document,
    ) -> None:
        section = document.sections[0]

        section.top_margin = self._cm(2.5)
        section.bottom_margin = self._cm(2.5)
        section.left_margin = self._cm(2.5)
        section.right_margin = self._cm(2.5)

        styles = document.styles

        normal_style = styles["Normal"]

        normal_style.font.name = self.FONT_NAME
        normal_style.font.size = Pt(self.FONT_SIZE)

        normal_style._element.rPr.rFonts.set(
            qn("w:ascii"),
            self.FONT_NAME,
        )

        normal_style._element.rPr.rFonts.set(
            qn("w:hAnsi"),
            self.FONT_NAME,
        )

        normal_style._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            self.FONT_NAME,
        )

        normal_style._element.rPr.rFonts.set(
            qn("w:cs"),
            self.FONT_NAME,
        )

    # ==================================================
    # Content
    # ==================================================

    def _add_heading(
        self,
        document: Document,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph()

        self._set_paragraph_rtl(paragraph)

        run = paragraph.add_run(text)

        self._set_run_font(
            run,
            size=self.HEADING_SIZE,
            bold=True,
        )

    def _add_content(
        self,
        document: Document,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph()

        self._set_paragraph_rtl(paragraph)

        self._set_paragraph_alignment(
            paragraph,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
        )

        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.5

        run = paragraph.add_run(text)

        self._set_run_font(
            run,
            size=self.FONT_SIZE,
        )

    def _add_faq_question(
        self,
        document: Document,
        question: str,
    ) -> None:
        paragraph = document.add_paragraph()

        self._set_paragraph_rtl(paragraph)

        run = paragraph.add_run(f"سوال: {question}")

        self._set_run_font(
            run,
            size=self.FONT_SIZE,
            bold=True,
        )

    # ==================================================
    # RTL
    # ==================================================

    def _set_paragraph_rtl(
        self,
        paragraph,
    ) -> None:
        paragraph_format = paragraph.paragraph_format

        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p_pr = paragraph._p.get_or_add_pPr()

        bidi = p_pr.find(qn("w:bidi"))

        if bidi is None:
            bidi = OxmlElement("w:bidi")
            p_pr.append(bidi)

        bidi.set(
            qn("w:val"),
            "1",
        )

    def _set_paragraph_alignment(
        self,
        paragraph,
        alignment,
    ) -> None:
        paragraph.alignment = alignment

    # ==================================================
    # Font
    # ==================================================

    def _set_run_font(
        self,
        run,
        size: int,
        bold: bool = False,
    ) -> None:
        run.font.name = self.FONT_NAME
        run.font.size = Pt(size)
        run.bold = bold

        r_pr = run._element.get_or_add_rPr()

        r_fonts = r_pr.rFonts

        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)

        r_fonts.set(
            qn("w:ascii"),
            self.FONT_NAME,
        )

        r_fonts.set(
            qn("w:hAnsi"),
            self.FONT_NAME,
        )

        r_fonts.set(
            qn("w:eastAsia"),
            self.FONT_NAME,
        )

        r_fonts.set(
            qn("w:cs"),
            self.FONT_NAME,
        )

    # ==================================================
    # Utilities
    # ==================================================

    def _cm(self, value: float):
        from docx.shared import Cm

        return Cm(value)
